from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from lead_data_validator import (
    check_geo,
    check_nwc,
    check_prooflink,
    check_title,
    parse_req,
    run_validation,
    validate,
)


def make_row(**overrides):
    row = {
        "first_name": "Jane",
        "last_name": "Doe",
        "company": "Acme",
        "title": "Director of Procurement",
        "prooflink": "linkedin.com/in/jane-doe",
        "location": "New York, USA",
        "status": "",
        "email": "jane@acme.com",
        "employees": 1200,
        "employees_prooflink": "www.acme.com/team",
        "industry": "Software",
        "req": "Job_level: Director + | Keywords: Procurement",
        "sub status": "N/A: Title/PL Summary",
    }
    normalized_overrides = {}
    for key, value in overrides.items():
        normalized_key = "sub status" if key == "sub_status" else key
        normalized_overrides[normalized_key] = value
    row.update(normalized_overrides)
    return row


def test_parse_req_reads_pipe_delimited_pairs():
    req = "Geo: US | Job_level: Director + | Keywords: Procurement"
    parsed = parse_req(req)
    assert parsed["geo"] == "US"
    assert parsed["job_level"] == "Director +"
    assert parsed["keywords"] == "Procurement"


def test_check_title_uses_phrase_based_level_detection():
    row = make_row(title="Security Specialist")
    result, comment = check_title(
        row,
        {"job_level": "Manager +", "keywords": "Security"},
    )
    assert result == "INVALID"
    assert "manager+" in comment.lower()


def test_check_title_ignores_non_deterministic_placeholders():
    row = make_row(title="Director of Procurement")
    result, comment = check_title(
        row,
        {"job_level": "see comment", "keywords": "-"},
    )
    assert result == "VALID"
    assert comment == "No deterministic title restrictions"


def test_check_title_skips_mixed_job_level_requirements():
    row = make_row(title="Director of Quality Assurance")
    result, comment = check_title(
        row,
        {
            "job_level": (
                "ANY level: IT Security; IT Administrator, Engineer, Specialist+: "
                "General, Infrastructure, Network Systems/Systems; ANY IT, Data "
                "Protection, Privacy Architect, Manager+: ANY"
            ),
            "keywords": "-",
        },
    )
    assert result == "VALID"
    assert comment == "No deterministic title restrictions"


def test_check_prooflink_accepts_company_domain_without_scheme():
    result, comment = check_prooflink("acme.com/team", "jane@acme.com")
    assert result == "VALID"
    assert "official website" in comment


def test_check_nwc_maps_recheck_statuses():
    result, comment = check_nwc("no info")
    assert result == "RECHECK"
    assert "manual recheck" in comment.lower()


def test_check_geo_requires_resolvable_location():
    result, comment = check_geo(
        {"location": "Greater Lille Metropolitan Area"},
        {"geo": "France, Germany"},
    )
    assert result == "VALID"
    assert "france" in comment.lower()


def test_validate_marks_bad_data_sub_status_as_invalid():
    result, comment = validate(make_row(sub_status="N2: Out of Business/Bad data"))
    assert result == "INVALID"
    assert "bad data" in comment.lower()


def test_validate_routes_to_baseline_other_checks():
    result, comment = validate(make_row(sub_status="N/A: Other", email="person@gmail.com"))
    assert result == "INVALID"
    assert "public mailbox" in comment.lower()


def test_run_validation_writes_outputs_and_hyperlinks(tmp_path: Path):
    input_path = tmp_path / "input.xlsx"
    output_xlsx = tmp_path / "output.xlsx"
    output_csv = tmp_path / "output.csv"

    rows = [
        make_row(),
        make_row(
            title="Marketing Manager",
            req="Job_level: Manager + | Keywords: Marketing",
        ),
        make_row(
            sub_status="N1: NWC",
            status="r",
        ),
        make_row(
            sub_status="N/A: Other",
            email="person@gmail.com",
        ),
        make_row(
            sub_status="N2: Out of Business/Bad data",
        ),
        make_row(
            sub_status="N/A: Country/GEO",
            location="Greater Lille Metropolitan Area",
            req="Geo: France, Germany",
        ),
    ]

    pd.DataFrame(rows).to_excel(input_path, index=False)
    validated = run_validation(str(input_path), str(output_xlsx), str(output_csv))

    assert output_xlsx.exists()
    assert output_csv.exists()
    assert list(validated["Result"]) == [
        "VALID",
        "VALID",
        "RECHECK",
        "INVALID",
        "INVALID",
        "VALID",
    ]
    assert "Comment" in validated.columns

    wb = load_workbook(output_xlsx)
    ws = wb.active
    headers = {str(cell.value).strip().lower(): cell.column for cell in ws[1]}
    prooflink_cell = ws.cell(row=2, column=headers["prooflink"])
    employee_prooflink_cell = ws.cell(row=2, column=headers["employees_prooflink"])

    assert prooflink_cell.hyperlink is not None
    assert prooflink_cell.hyperlink.target == "https://linkedin.com/in/jane-doe"
    assert employee_prooflink_cell.hyperlink is not None
    assert employee_prooflink_cell.hyperlink.target == "https://www.acme.com/team"
