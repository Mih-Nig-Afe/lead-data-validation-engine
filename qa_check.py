#!/usr/bin/env python3
import argparse
from typing import Optional, Sequence

import pandas as pd
from openpyxl import load_workbook

from lead_data_validator import DEFAULT_INPUT_FILE, DEFAULT_OUTPUT_XLSX


ALLOWED_RESULTS = {"VALID", "INVALID", "RECHECK"}


def run_checks(input_file: str, output_file: str) -> pd.DataFrame:
    src = pd.read_excel(input_file, engine="openpyxl")
    out = pd.read_excel(output_file, engine="openpyxl")

    wb = load_workbook(output_file)
    ws = wb[wb.sheetnames[0]]

    assert len(src) == len(out), f"Row count mismatch: {len(src)} vs {len(out)}"
    assert "Result" in out.columns, "Missing Result column"
    assert "Comment" in out.columns, "Missing Comment column"

    result_col = len(src.columns) + 1
    comment_col = len(src.columns) + 2

    assert (
        ws.cell(row=1, column=result_col).value == "Result"
    ), "Result header is not appended after the source columns"
    assert (
        ws.cell(row=1, column=comment_col).value == "Comment"
    ), "Comment header is not appended after the source columns"

    bad = out[~out["Result"].astype(str).isin(ALLOWED_RESULTS)]
    assert bad.empty, f'Unexpected Result values: {bad["Result"].dropna().unique()[:10]}'

    empty_comments = out[out["Comment"].fillna("").astype(str).str.strip() == ""]
    assert empty_comments.empty, f"Empty comments found: {len(empty_comments)}"

    for row_idx in range(2, len(src) + 2):
        visible_result = str(ws.cell(row=row_idx, column=result_col).value or "").strip()
        visible_comment = str(
            ws.cell(row=row_idx, column=comment_col).value or ""
        ).strip()
        assert (
            visible_result in ALLOWED_RESULTS
        ), f"Invalid Result value at row {row_idx}: {visible_result}"
        assert visible_comment != "", f"Empty Comment cell at row {row_idx}"

    prooflink_headers = {
        str(cell.value).strip().lower(): cell.column
        for cell in ws[1]
        if cell.value is not None
    }
    for header in ("prooflink", "employees_prooflink"):
        col_idx = prooflink_headers.get(header)
        if not col_idx:
            continue
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value:
                assert cell.hyperlink is not None, (
                    f"Expected hyperlink in {header} at row {row_idx}"
                )

    n1 = out[out["sub status"].astype(str).str.strip().str.lower() == "n1: nwc"]
    for index, row in n1.iterrows():
        status = str(row.get("status", "")).strip().lower()
        result = str(row.get("Result", "")).strip()
        if status in ("", "valid"):
            assert result == "VALID", f"N1 expected VALID at row {index + 2}"
        elif status == "a":
            assert result == "INVALID", f"N1 expected INVALID at row {index + 2}"
        elif status == "!":
            assert result == "INVALID", f"N1 expected INVALID at row {index + 2}"
        elif status in ("r", "no info", "no company match"):
            assert result == "RECHECK", f"N1 expected RECHECK at row {index + 2}"

    return out


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Run QA checks against the generated validation workbook."
    )
    parser.add_argument(
        "--input",
        default=DEFAULT_INPUT_FILE,
        help=f"Original source workbook path. Default: {DEFAULT_INPUT_FILE}",
    )
    parser.add_argument(
        "--output",
        default=DEFAULT_OUTPUT_XLSX,
        help=f"Generated workbook path. Default: {DEFAULT_OUTPUT_XLSX}",
    )
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    validated = run_checks(args.input, args.output)
    print("QA OK")
    print("Rows:", len(validated))
    print("Result distribution:", validated["Result"].value_counts(dropna=False).to_dict())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
