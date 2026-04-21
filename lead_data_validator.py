#!/usr/bin/env python3
import argparse
import re
from typing import Dict, List, Mapping, Optional, Sequence, Tuple
from urllib.parse import urlparse

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


PUBLIC_EMAIL_DOMAINS = {
    "gmail.com",
    "yahoo.com",
    "hotmail.com",
    "outlook.com",
    "aol.com",
    "icloud.com",
    "proton.me",
    "protonmail.com",
    "mail.com",
    "gmx.com",
}

EMAIL_REGEX = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
WORD_REGEX = re.compile(r"\w+")
PLACEHOLDER_REQ_VALUES = {"", "-", "any", "n/a", "na"}
NON_DETERMINISTIC_REQ_MARKERS = ("see comment", "see comments")
NON_DETERMINISTIC_LEVEL_MARKERS = ("any level", "all job levels", "see comment")
LEVEL_NAME_BY_RANK = {
    1: "individual contributor",
    2: "lead",
    3: "manager",
    4: "director",
    5: "vp",
    6: "c-level",
}
LEVEL_PATTERNS = (
    (
        6,
        (
            r"\bc[\s-]?suite\b",
            r"\bc[\s-]?level\b",
            r"\bchief\b",
            r"\bceo\b",
            r"\bcfo\b",
            r"\bcio\b",
            r"\bcoo\b",
            r"\bcto\b",
            r"\bcmo\b",
            r"\bchro\b",
            r"\bpresident\b",
        ),
    ),
    (
        5,
        (
            r"\bexecutive vice president\b",
            r"\bsenior vice president\b",
            r"\bvice president\b",
            r"\bevp\b",
            r"\bsvp\b",
            r"\bavp\b",
            r"\bvp\b",
        ),
    ),
    (4, (r"\bsenior director\b", r"\bdirector\b")),
    (3, (r"\bsenior manager\b", r"\bmanager\b", r"\bhead\b")),
    (2, (r"\blead\b", r"\bprincipal\b")),
    (
        1,
        (
            r"\barchitect\b",
            r"\bengineer\b",
            r"\bspecialist\b",
            r"\badministrator\b",
            r"\banalyst\b",
            r"\bcoordinator\b",
            r"\bassociate\b",
            r"\bemployee\b",
            r"\bindividual contributor\b",
        ),
    ),
)

DEFAULT_INPUT_FILE = "DataCheck_DemoCode.xlsx"
DEFAULT_OUTPUT_XLSX = "Lead_Data_Validation_Results.xlsx"
DEFAULT_OUTPUT_CSV = "Lead_Data_Validation_Results.csv"
COUNTRY_ALIASES = {
    "australia": "australia",
    "austria": "austria",
    "bahrain": "bahrain",
    "belgium": "belgium",
    "canada": "canada",
    "denmark": "denmark",
    "finland": "finland",
    "france": "france",
    "germany": "germany",
    "hyderabad": "india",
    "india": "india",
    "ireland": "ireland",
    "italy": "italy",
    "kuwait": "kuwait",
    "liechtenstein": "liechtenstein",
    "lille": "france",
    "melbourne": "australia",
    "mumbai": "india",
    "netherlands": "netherlands",
    "oman": "oman",
    "portugal": "portugal",
    "qatar": "qatar",
    "saudi arabia": "saudi arabia",
    "slovakia": "slovakia",
    "south africa": "south africa",
    "spain": "spain",
    "sweden": "sweden",
    "switzerland": "switzerland",
    "sydney": "australia",
    "turkey": "turkey",
    "uae": "united arab emirates",
    "uk": "united kingdom",
    "united arab emirates": "united arab emirates",
    "united kingdom": "united kingdom",
    "us": "united states",
    "usa": "united states",
    "united states": "united states",
}


def norm(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


def parse_req(req: str) -> Dict[str, str]:
    out: Dict[str, str] = {}
    for part in norm(req).split("|"):
        if ":" not in part:
            continue
        key, value = part.split(":", 1)
        out[key.strip().lower()] = value.strip()
    return out


def normalize_requirement_text(value: str) -> str:
    return re.sub(r"\s+", " ", norm(value).lower()).strip()


def requirement_is_placeholder(value: str) -> bool:
    normalized = normalize_requirement_text(value)
    if normalized in PLACEHOLDER_REQ_VALUES:
        return True
    return any(marker in normalized for marker in NON_DETERMINISTIC_REQ_MARKERS)


def email_domain(email: str) -> str:
    return email.lower().split("@")[-1] if "@" in email else ""


def base_domain(domain: str) -> str:
    domain = domain.replace("www.", "")
    parts = domain.split(".")
    return ".".join(parts[-2:]) if len(parts) >= 2 else domain


def extract_level_ranks(text: str) -> List[int]:
    normalized = normalize_requirement_text(text)
    if not normalized:
        return []

    ranks = []
    for rank, patterns in LEVEL_PATTERNS:
        if any(re.search(pattern, normalized) for pattern in patterns):
            ranks.append(rank)
    return ranks


def required_title_level(level_requirement: str) -> int:
    if requirement_is_placeholder(level_requirement):
        return 0

    normalized = normalize_requirement_text(level_requirement)
    if any(marker in normalized for marker in NON_DETERMINISTIC_LEVEL_MARKERS):
        return 0

    ranks = extract_level_ranks(level_requirement)
    if len(set(ranks)) > 1:
        return 0
    return min(ranks) if ranks else 0


def inferred_title_level(title: str) -> int:
    ranks = extract_level_ranks(title)
    return max(ranks) if ranks else 1


def extract_keyword_candidates(raw_keywords: str) -> List[str]:
    if requirement_is_placeholder(raw_keywords):
        return []

    candidates = []
    for raw_candidate in re.split(r"[;,]", norm(raw_keywords)):
        candidate = raw_candidate.strip()
        if not candidate:
            continue
        candidate = re.sub(r".*keywords:\s*", "", candidate, flags=re.IGNORECASE)
        candidate = candidate.strip(" .-")
        if not candidate:
            continue
        if requirement_is_placeholder(candidate):
            continue
        if not WORD_REGEX.findall(candidate):
            continue
        candidates.append(candidate)
    return candidates


def extract_allowed_countries(raw_geo: str) -> List[str]:
    if requirement_is_placeholder(raw_geo):
        return []

    normalized = normalize_requirement_text(raw_geo)
    normalized = re.sub(r"\bonly\b", " ", normalized)
    tokens = re.split(r"[,:;/&()\-]| and ", normalized)
    countries = []
    seen = set()
    for token in tokens:
        token = token.strip()
        if not token:
            continue
        country = COUNTRY_ALIASES.get(token, token if token in COUNTRY_ALIASES.values() else None)
        if not country or country in seen:
            continue
        seen.add(country)
        countries.append(country)
    return countries


def infer_country_from_location(location: str) -> Optional[str]:
    normalized = normalize_requirement_text(location)
    if not normalized:
        return None

    matches = []
    for token, country in COUNTRY_ALIASES.items():
        if re.search(rf"\b{re.escape(token)}\b", normalized):
            matches.append(country)

    unique_matches = list(dict.fromkeys(matches))
    if len(unique_matches) == 1:
        return unique_matches[0]
    return None


def normalize_url_for_hyperlink(raw_link: str) -> Optional[str]:
    link = norm(raw_link)
    if not link:
        return None

    parsed = urlparse(link)
    if parsed.scheme in {"http", "https"} and parsed.netloc:
        return link

    candidate = f"https://{link}"
    parsed_candidate = urlparse(candidate)
    if parsed_candidate.netloc and "." in parsed_candidate.netloc:
        return candidate
    return None


def check_prooflink(link: str, email: str) -> Tuple[str, str]:
    normalized_link = norm(link).lower()
    normalized_email = norm(email)

    if not normalized_link:
        return "INVALID", "Missing prooflink"

    if "linkedin.com/in/" in normalized_link or "zoominfo.com/p/" in normalized_link:
        return "VALID", "Valid prooflink: trusted source (LinkedIn/ZoomInfo)"

    parsed_link = normalize_url_for_hyperlink(normalized_link) or normalized_link
    host = urlparse(parsed_link).netloc.replace("www.", "")
    if host and base_domain(host) == base_domain(email_domain(normalized_email)):
        return (
            "VALID",
            "Valid prooflink: official website matches corporate email domain",
        )

    return (
        "INVALID",
        "Invalid prooflink source: expected linkedin.com/in/, zoominfo.com/p/, or company domain match",
    )


def check_geo(row: Mapping[str, object], req_map: Mapping[str, str]) -> Tuple[str, str]:
    allowed_countries = extract_allowed_countries(req_map.get("geo", ""))
    if not allowed_countries:
        return "RECHECK", "Manual GEO review needed: requirement is not machine-readable"

    location = norm(row.get("location"))
    if not location:
        return "INVALID", "Missing location"

    inferred_country = infer_country_from_location(location)
    if not inferred_country:
        return "RECHECK", "Manual GEO review needed: location does not resolve to a single country"

    if inferred_country in allowed_countries:
        return "VALID", f"Location resolves to allowed country: {inferred_country.title()}"
    return "INVALID", f"Location resolves outside GEO requirement: {inferred_country.title()}"


def check_bad_data_status(sub_status: str) -> Tuple[str, str]:
    normalized = normalize_requirement_text(sub_status)
    if "company" in normalized:
        return "INVALID", "Company marked as out of business or bad data"
    return "INVALID", "Lead marked as out of business or bad data"


def check_title(row: Mapping[str, object], req_map: Mapping[str, str]) -> Tuple[str, str]:
    title = norm(row.get("title")).lower()
    if not title:
        return "INVALID", "Missing title"

    issues = []
    matched_keyword = None

    keyword_candidates = extract_keyword_candidates(req_map.get("keywords", ""))
    if keyword_candidates:
        title_tokens = set(WORD_REGEX.findall(title))
        for keyword in keyword_candidates:
            keyword_tokens = set(WORD_REGEX.findall(keyword.lower()))
            if keyword_tokens and keyword_tokens.issubset(title_tokens):
                matched_keyword = keyword
                break
        if not matched_keyword:
            issues.append("Title missing required keywords")

    req_level = required_title_level(req_map.get("job_level", ""))
    title_level = inferred_title_level(title)
    if req_level and title_level < req_level:
        issues.append(
            f"Title level below requirement ({LEVEL_NAME_BY_RANK[req_level]}+ expected)"
        )

    if issues:
        return "INVALID", "; ".join(issues)

    positive_notes = []
    if matched_keyword:
        positive_notes.append(f"Title contains keywords: {matched_keyword}")
    if req_level:
        positive_notes.append(
            f"Title level meets minimum: {LEVEL_NAME_BY_RANK[req_level]}+"
        )

    if positive_notes:
        return "VALID", "; ".join(positive_notes)
    return "VALID", "No deterministic title restrictions"


def check_nwc(status: str) -> Tuple[str, str]:
    normalized_status = norm(status).lower()
    if not normalized_status or normalized_status == "valid":
        return "VALID", "NWC status clear"
    if normalized_status == "a":
        return "INVALID", "Retrieved lead"
    if normalized_status == "!":
        return "INVALID", "Suspicious lead"
    if normalized_status in {"r", "no info", "no company match"}:
        return "RECHECK", "Manual recheck needed"
    return "INVALID", f"Unknown NWC status: {normalized_status}"


def check_other(row: Mapping[str, object]) -> Tuple[str, str]:
    needed = ["first_name", "last_name", "company", "email"]
    missing = [column for column in needed if not norm(row.get(column))]
    if missing:
        return "INVALID", f"Missing required data: {', '.join(missing)}"

    email = norm(row.get("email")).lower()
    if not EMAIL_REGEX.match(email):
        return "INVALID", "Invalid email format"
    if email_domain(email) in PUBLIC_EMAIL_DOMAINS:
        return "INVALID", "Email must be corporate (public mailbox detected)"

    return "VALID", "Lead data matches requirements"


def validate(row: Mapping[str, object]) -> Tuple[str, str]:
    sub_status = norm(row.get("sub status")).lower()
    req_map = parse_req(norm(row.get("req")))

    if "nwc" in sub_status:
        return check_nwc(row.get("status"))
    if "country" in sub_status or "geo" in sub_status:
        return check_geo(row, req_map)
    if "prooflink" in sub_status:
        return check_prooflink(row.get("prooflink"), row.get("email"))
    if "title" in sub_status:
        return check_title(row, req_map)
    if "out of business" in sub_status or "bad data" in sub_status:
        return check_bad_data_status(sub_status)
    if "other" in sub_status:
        return check_other(row)

    return check_other(row)


def validate_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    results: List[str] = []
    comments: List[str] = []

    for row in df.to_dict(orient="records"):
        result, comment = validate(row)
        results.append(result)
        comments.append(comment)

    enriched = df.copy()
    enriched["Result"] = results
    enriched["Comment"] = comments
    return enriched


def format_output_workbook(workbook_path: str) -> None:
    wb = load_workbook(workbook_path)
    ws = wb.active

    hyperlink_columns = set()
    wrapped_headers = {"result", "comment", "req"}
    wrapped_columns = set()

    for idx, cell in enumerate(ws[1], start=1):
        header = norm(cell.value).lower()
        if header in {"prooflink", "employees_prooflink"}:
            hyperlink_columns.add(idx)
        if header in wrapped_headers:
            wrapped_columns.add(idx)

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)

        for cell in col:
            cell_value = str(cell.value) if cell.value else ""
            max_len = max(max_len, len(cell_value))

            if cell.column in wrapped_columns:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="center")

            if cell.column in hyperlink_columns and cell.row > 1 and cell.value:
                target = normalize_url_for_hyperlink(str(cell.value))
                if target:
                    cell.hyperlink = target
                    cell.style = "Hyperlink"

        ws.column_dimensions[col_letter].width = min(max_len + 2, 45)

    wb.save(workbook_path)


def run_validation(
    input_file: str = DEFAULT_INPUT_FILE,
    output_xlsx: str = DEFAULT_OUTPUT_XLSX,
    output_csv: str = DEFAULT_OUTPUT_CSV,
) -> pd.DataFrame:
    df = pd.read_excel(input_file, engine="openpyxl")
    validated = validate_dataframe(df)
    validated.to_excel(output_xlsx, index=False)
    validated.to_csv(output_csv, index=False)
    format_output_workbook(output_xlsx)
    return validated


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Validate lead records from an input workbook and write XLSX/CSV outputs."
    )
    parser.add_argument(
        "--input",
        default=DEFAULT_INPUT_FILE,
        help=f"Input workbook path. Default: {DEFAULT_INPUT_FILE}",
    )
    parser.add_argument(
        "--output-xlsx",
        default=DEFAULT_OUTPUT_XLSX,
        help=f"Formatted workbook output path. Default: {DEFAULT_OUTPUT_XLSX}",
    )
    parser.add_argument(
        "--output-csv",
        default=DEFAULT_OUTPUT_CSV,
        help=f"CSV output path. Default: {DEFAULT_OUTPUT_CSV}",
    )
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    run_validation(args.input, args.output_xlsx, args.output_csv)
    print("Done:", args.output_xlsx, args.output_csv)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
